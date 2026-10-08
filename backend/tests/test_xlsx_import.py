"""Excel(.xlsx) 批量导入测试：解析服务 + /api/products/parse-xlsx 端点。"""
import base64
import io

import pytest
from openpyxl import Workbook

from app.services import xlsx_service

BASE = "/api/products"


def make_xlsx(rows, sheet_title="产品"):
    """用 openpyxl 现造一个 .xlsx 并返回 base64。"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return base64.b64encode(buf.getvalue()).decode()


HEADER_SUMMER = {
    "filename": "products.xlsx",
    "content_base64": make_xlsx(
        [
            ["名称", "型号", "分类", "参数", "应用场景"],
            ["便携式储能电源", "AF-PS600", "储能", "容量: 600Wh; 输出: 600W", "露营; 应急备电"],
            ["太阳能折叠板", "AF-SP120", "光伏", "功率: 120W", "房车供电"],
        ]
    ),
}


# ---------------- 纯函数层 ----------------

def test_parse_rows_skips_header():
    rows = [["名称", "型号", "分类", "参数", "应用场景"], ["储能电源", "AF-PS600", "储能", "", ""]]
    out = xlsx_service.parse_xlsx_rows(rows)
    assert out["skippedHeader"] is True
    assert len(out["items"]) == 1
    assert out["items"][0]["name"] == "储能电源"


def test_parse_rows_without_header():
    rows = [["储能电源", "AF-PS600", "储能", "", ""]]
    out = xlsx_service.parse_xlsx_rows(rows)
    assert out["skippedHeader"] is False
    assert len(out["items"]) == 1


def test_parse_rows_splits_multi_value_on_common_separators():
    rows = [["产品A", "M1", "类", "容量: 1kWh；输出: 5kW; 重量: 30kg", "户外｜工地; 应急"]]
    out = xlsx_service.parse_xlsx_rows(rows)
    item = out["items"][0]
    assert "容量: 1kWh" in item["specs"]
    assert "输出: 5kW" in item["specs"]
    assert "重量: 30kg" in item["specs"]
    assert item["applications"] == ["户外", "工地", "应急"]


def test_parse_rows_empty_name_reported():
    rows = [["名称", "型号"], ["", "AF-X"], ["有效产品", "AF-Y"]]
    out = xlsx_service.parse_xlsx_rows(rows)
    assert len(out["items"]) == 1
    assert out["items"][0]["name"] == "有效产品"
    assert any("名称为空" in e for e in out["errors"])


def test_parse_rows_empty_sheet():
    out = xlsx_service.parse_xlsx_rows([])
    assert out["items"] == []
    assert out["errors"]


def test_parse_rows_blank_rows_ignored():
    rows = [["名称"], ["  ", "  "], ["真实产品"]]
    out = xlsx_service.parse_xlsx_rows(rows)
    assert [i["name"] for i in out["items"]] == ["真实产品"]


def test_parse_rows_caps_at_200():
    rows = [["名称", "型号"]] + [[f"产品{i}", "M"] for i in range(250)]
    out = xlsx_service.parse_xlsx_rows(rows)
    assert len(out["items"]) == 200
    assert any("上限" in e for e in out["errors"])


def test_numeric_cells_are_readable():
    """Excel 里的数字应渲染成 600 而不是 600.0。"""
    out = xlsx_service.parse_xlsx_rows([["储能", 600, "储能", 600.0, ""]])
    assert out["items"][0]["model"] == "600"
    assert out["items"][0]["category"] == "储能"


# ---------------- base64 / 端点层 ----------------

def test_parse_xlsx_base64_real_file():
    out = xlsx_service.parse_xlsx_base64(HEADER_SUMMER["content_base64"])
    assert out["skippedHeader"] is True
    assert len(out["items"]) == 2
    assert out["items"][0]["name"] == "便携式储能电源"
    assert out["sheet"] == "产品"


def test_parse_xlsx_rejects_non_xlsx_bytes():
    junk = base64.b64encode(b"this is not a spreadsheet").decode()
    with pytest.raises(xlsx_service.XlsxParseError):
        xlsx_service.parse_xlsx_base64(junk)


def test_parse_xlsx_rejects_invalid_base64():
    with pytest.raises(xlsx_service.XlsxParseError):
        xlsx_service.parse_xlsx_base64("!!!not-base64!!!")


def test_parse_xlsx_rejects_empty():
    with pytest.raises(xlsx_service.XlsxParseError):
        xlsx_service.parse_xlsx_base64("")


def test_parse_xlsx_accepts_data_uri_prefix():
    uri = "data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64," + HEADER_SUMMER["content_base64"]
    out = xlsx_service.parse_xlsx_base64(uri)
    assert len(out["items"]) == 2


def test_endpoint_parse_only_does_not_persist(client):
    """commit=false：只预览，不应写库。"""
    before = len(client.get(BASE).json())
    r = client.post(f"{BASE}/parse-xlsx", json=HEADER_SUMMER)
    assert r.status_code == 200
    d = r.json()
    assert d["count"] == 2
    assert d["skippedHeader"] is True
    assert len(d["items"]) == 2
    assert "created" not in d
    assert len(client.get(BASE).json()) == before


def test_endpoint_parse_and_commit_persists(client):
    r = client.post(f"{BASE}/parse-xlsx", json={**HEADER_SUMMER, "commit": True})
    assert r.status_code == 200
    d = r.json()
    assert d["created_count"] == 2
    assert d["failed_count"] == 0

    names = [p["name"] for p in client.get(BASE).json()]
    assert "便携式储能电源" in names
    assert "太阳能折叠板" in names


def test_endpoint_commit_uses_provided_items(client):
    """commit 时可回传人工编辑过的 items，应以回传内容为准。"""
    payload = {
        **HEADER_SUMMER,
        "commit": True,
        "items": [{"name": "人工改名后的产品", "model": "AF-EDIT", "category": "测试"}],
    }
    r = client.post(f"{BASE}/parse-xlsx", json=payload)
    assert r.status_code == 200
    assert r.json()["created_count"] == 1

    names = [p["name"] for p in client.get(BASE).json()]
    assert "人工改名后的产品" in names
    assert "便携式储能电源" not in names


def test_endpoint_invalid_file_400(client):
    r = client.post(f"{BASE}/parse-xlsx", json={"filename": "x.xlsx", "content_base64": base64.b64encode(b"junk").decode()})
    assert r.status_code == 400
    assert "xlsx" in r.json()["detail"].lower() or "Excel" in r.json()["detail"]


def test_endpoint_missing_content_422(client):
    assert client.post(f"{BASE}/parse-xlsx", json={"filename": "x.xlsx"}).status_code == 422


def test_endpoint_multisheet_uses_first(client):
    wb = Workbook()
    ws = wb.active
    ws.title = "第一个表"
    ws.append(["名称", "型号"])
    ws.append(["表一产品", "AF-1"])
    ws2 = wb.create_sheet("第二个表")
    ws2.append(["名称", "型号"])
    ws2.append(["表二产品", "AF-2"])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    b64 = base64.b64encode(buf.getvalue()).decode()

    r = client.post(f"{BASE}/parse-xlsx", json={"filename": "m.xlsx", "content_base64": b64})
    assert r.status_code == 200
    d = r.json()
    assert d["sheet"] == "第一个表"
    assert d["items"][0]["name"] == "表一产品"


def test_parse_xlsx_row_with_colon_spec_becomes_pair(client):
    """导入后参数应被拆成 {key, value}。"""
    client.post(f"{BASE}/parse-xlsx", json={**HEADER_SUMMER, "commit": True})
    product = next(p for p in client.get(BASE).json() if p["name"] == "便携式储能电源")
    kv = {s["key"]: s["value"] for s in product["specs"]}
    assert kv["容量"] == "600Wh"
    assert kv["输出"] == "600W"
